"""
Generate a DOCX evaluation form (Julian form) from interview transcript and Gemini notes.
Uses Gemini to answer the structured questions and python-docx to build the document.
"""
import os
import re
from pathlib import Path

_config_env = Path(__file__).resolve().parent.parent / "config" / ".env"
if _config_env.exists():
    from dotenv import load_dotenv
    load_dotenv(_config_env)

RECOMMENDATION_OPTIONS = [
    "Strongly Recommend",
    "Recommend",
    "Neutral",
    "Do Not Recommend",
    "Strongly Do Not Recommend",
]


def _extract_answers_with_gemini(transcript_text: str, candidate_name: str = "", *, model: str | None = None, api_key: str | None = None) -> dict:
    """
    Use Gemini to answer the Julian form questions from the transcript.
    Returns dict: overall_impression (1-10), key_strengths, areas_for_development, recommendation (exactly one of RECOMMENDATION_OPTIONS), additional_comments.
    """
    from google import genai

    model = model or os.getenv("GEMINI_MODEL", "gemini-2.5-flash")
    api_key = api_key or os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise ValueError("GEMINI_API_KEY is not set")

    client = genai.Client(api_key=api_key)
    opts = " | ".join(RECOMMENDATION_OPTIONS)

    prompt = f"""You are an expert interviewer. Based on the following interview transcript (including any Gemini notes or evaluator notes), fill in the evaluation form. Use the **full** transcript.

Candidate: {candidate_name or "Candidate"}

## Interview transcript (use everything below)
{transcript_text}

---
Reply with exactly this format (no other text before or after). Use a single line for each field; for multi-line answers use \\n for newlines.

OVERALL_IMPRESSION: <integer from 1 to 10, where 1=Poor and 10=Excellent>
KEY_STRENGTHS: <2-5 bullet points or short paragraphs on the candidate's key strengths based on the interview>
AREAS_FOR_DEVELOPMENT: <2-4 bullet points or short paragraphs on areas for development>
RECOMMENDATION: <exactly one of: {opts}>
ADDITIONAL_COMMENTS: <any other observations or comments; can be "None" if nothing to add>
"""

    response = client.models.generate_content(model=model, contents=prompt)
    text = (response.text or "").strip()

    result = {
        "overall_impression": 5,
        "key_strengths": "",
        "areas_for_development": "",
        "recommendation": "Neutral",
        "additional_comments": "",
    }
    # Parse key: value, with optional multi-line values (until next KEY: or end)
    keys = ["OVERALL_IMPRESSION", "KEY_STRENGTHS", "AREAS_FOR_DEVELOPMENT", "RECOMMENDATION", "ADDITIONAL_COMMENTS"]
    pattern = re.compile(r"^(" + "|".join(re.escape(k) for k in keys) + r")\s*:\s*(.*)$", re.IGNORECASE)
    current_key = None
    current_val = []
    lines = text.split("\n")
    for i, line in enumerate(lines):
        m = pattern.match(line.strip())
        if m:
            if current_key:
                val = "\n".join(current_val).strip().replace("\\n", "\n")
                _apply_parsed(result, current_key, val)
            current_key = m.group(1).upper()
            current_val = [m.group(2).strip()]
        elif current_key:
            current_val.append(line.strip())
    if current_key:
        val = "\n".join(current_val).strip().replace("\\n", "\n")
        _apply_parsed(result, current_key, val)

    result["overall_impression"] = max(1, min(10, result["overall_impression"]))
    if result["recommendation"] not in RECOMMENDATION_OPTIONS:
        for opt in RECOMMENDATION_OPTIONS:
            if opt.lower() in (result["recommendation"] or "").lower():
                result["recommendation"] = opt
                break
        else:
            result["recommendation"] = "Neutral"
    return result


def _apply_parsed(result: dict, key: str, val: str) -> None:
    if key == "OVERALL_IMPRESSION":
        m = re.search(r"[1-9]|10", val)
        if m:
            result["overall_impression"] = int(m.group())
    elif key == "KEY_STRENGTHS":
        result["key_strengths"] = val
    elif key == "AREAS_FOR_DEVELOPMENT":
        result["areas_for_development"] = val
    elif key == "RECOMMENDATION":
        result["recommendation"] = val
    elif key == "ADDITIONAL_COMMENTS":
        result["additional_comments"] = val


def build_julian_docx(
    transcript_text: str,
    candidate_name: str = "",
    output_path: Path | str | None = None,
    *,
    model: str | None = None,
    api_key: str | None = None,
) -> Path:
    """
    Generate the Julian form DOCX from transcript and save to output_path.
    If output_path is None, saves to data/results/julian_form.docx (caller can pass a path with candidate email).
    Returns the path to the saved file.
    """
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    answers = _extract_answers_with_gemini(transcript_text, candidate_name=candidate_name, model=model, api_key=api_key)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.size = Pt(11)

    # Title
    doc.add_paragraph("Interview Evaluation Form (Julian)", style="Heading 1")
    if candidate_name:
        doc.add_paragraph(f"Candidate: {candidate_name}")
    doc.add_paragraph()

    # Overall Impression
    doc.add_paragraph("Overall Impression of the Candidate (based on your interview topic)", style="Heading 2")
    scale = "Poor  1  2  3  4  5  6  7  8  9  10  Excellent"
    p = doc.add_paragraph(scale)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Score: {answers['overall_impression']}")
    doc.add_paragraph()

    # Key strengths
    doc.add_paragraph("What are the candidate's key strengths? (Based on your interview topic)", style="Heading 2")
    doc.add_paragraph(answers["key_strengths"] or "—")
    doc.add_paragraph()

    # Areas for development
    doc.add_paragraph("What are the candidate's areas for development? (Based on your interview topic)", style="Heading 2")
    doc.add_paragraph(answers["areas_for_development"] or "—")
    doc.add_paragraph()

    # Recommendation
    doc.add_paragraph(
        "Would you recommend this candidate for the position (overall impression, can go outside of your interview topics)",
        style="Heading 2",
    )
    for opt in RECOMMENDATION_OPTIONS:
        doc.add_paragraph(opt)
    doc.add_paragraph(f"Selected: {answers['recommendation']}")
    doc.add_paragraph()

    # Additional comments
    doc.add_paragraph("Any additional comments or observations?", style="Heading 2")
    doc.add_paragraph(answers["additional_comments"] or "—")

    out = Path(output_path) if output_path else Path(__file__).resolve().parent.parent / "data" / "results" / "julian_form.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
